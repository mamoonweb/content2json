#!/usr/bin/env python3
"""
content_to_json.py - Map a Word content file onto an Elementor JSON template.

Commands
  inspect  --content X.docx --template T.json [--mapping M.json]
           Show the sections found in the content file and the widgets (with ids)
           found in the template. Use it to author a mapping config.
  map      --content X.docx --template T.json --mapping M.json --out OUT.json
           [--report REPORT.md] [--strict]
           Populate the template, validate, write the JSON + reports.

The engine contains NO page-specific logic. All mapping lives in the mapping
config (JSON), so another content file with the same layout can be processed by
re-using the config; a content file with a different layout needs only a new
config.
"""
import argparse
import copy
import html
import json
import re
import secrets
import sys
from dataclasses import dataclass, field
from typing import Any, List, Optional

# ----------------------------------------------------------------------------
# helpers
# ----------------------------------------------------------------------------
TAG_RE = re.compile(r"<[^>]+>")


def strip_tags(s: str) -> str:
    return TAG_RE.sub(" ", s or "")


def norm(s: str) -> str:
    s = html.unescape(strip_tags(s)).replace("\u00a0", " ")
    return re.sub(r"\s+", " ", s).strip()


def esc(s: str) -> str:
    return html.escape(s, quote=False)


def excerpt(s: Any, n: int = 70) -> str:
    s = norm(s) if isinstance(s, str) else json.dumps(s, ensure_ascii=False)
    return s if len(s) <= n else s[: n - 1] + "…"


# ----------------------------------------------------------------------------
# source model
# ----------------------------------------------------------------------------
CONTENT_KINDS = ("heading", "paragraph", "button", "listitem", "label")


@dataclass
class Node:
    uid: int
    kind: str  # heading|paragraph|button|listitem|label|placeholder
    text: str
    level: int = 0
    bold: bool = False
    cell: int = 0
    used: list = field(default_factory=list)  # [(rule, target, field)]

    def short(self):
        return excerpt(self.text)


@dataclass
class Card:
    heading: Node
    body: List[Node] = field(default_factory=list)
    buttons: List[Node] = field(default_factory=list)
    placeholders: List[Node] = field(default_factory=list)
    pre_placeholders: List[Node] = field(default_factory=list)


@dataclass
class Section:
    index: int
    label: str
    nodes: List[Node]
    title: Optional[Node] = None
    intro: List[Node] = field(default_factory=list)
    outro: List[Node] = field(default_factory=list)
    cards: List[Card] = field(default_factory=list)
    buttons: List[Node] = field(default_factory=list)
    placeholders: List[Node] = field(default_factory=list)
    items: List[Node] = field(default_factory=list)
    rule: Optional[str] = None
    ignored: Optional[str] = None


DEFAULT_SRC_OPTS = {
    "button_max_words": 4,
    "button_suffix": " Button",
    "placeholder_patterns": [r"^\*[^*].*\*$", r"^\s*insert\b", r"^\[\s*social media"],
    "bullet_chars": "·•▪‣-–",
}


class SourceDoc:
    def __init__(self):
        self.title = ""
        self.sections: List[Section] = []
        self.notes: List[str] = []  # header/footer text, not page content


def parse_docx(path: str, opts: dict) -> SourceDoc:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    o = {**DEFAULT_SRC_OPTS, **(opts or {})}
    ph_res = [re.compile(p, re.I) for p in o["placeholder_patterns"]]
    document = Document(path)
    doc = SourceDoc()
    counters = {"uid": 0, "cell": 0}

    def para_text(p):
        out = []
        for el in p.iter():
            if el.tag == qn("w:t"):
                out.append(el.text or "")
            elif el.tag in (qn("w:br"), qn("w:cr")):
                out.append("\n")
            elif el.tag == qn("w:tab"):
                out.append(" ")
        t = "".join(out)
        t = re.sub(r"[ \t\u00a0]+", " ", t)
        t = re.sub(r" *\n *", "\n", t)
        return t.strip()

    def para_bold(p):
        seen = False
        for r in p.iter(qn("w:r")):
            txt = "".join(t.text or "" for t in r.findall(qn("w:t")))
            if not txt.strip():
                continue
            b = r.find(qn("w:rPr") + "/" + qn("w:b"))
            if b is None or b.get(qn("w:val")) in ("0", "false"):
                return False
            seen = True
        return seen

    def make_node(p, cell):
        text = para_text(p)
        if not text:
            return None
        style = Paragraph(p, document).style.name or ""
        m = re.match(r"(?i)heading\s*(\d)", style)
        bold = para_bold(p)
        counters["uid"] += 1
        uid = counters["uid"]
        if any(r.search(text) for r in ph_res):
            return Node(uid, "placeholder", text, 0, bold, cell)
        if m:
            return Node(uid, "heading", text, int(m.group(1)), bold, cell)
        if text[0] in o["bullet_chars"] and len(text) > 1 and text[1] in " \t":
            return Node(uid, "listitem", text[1:].strip(), 0, bold, cell)
        mb = re.match(r"^\[(.+)\]$", text)
        if mb:
            return Node(uid, "button", mb.group(1).strip(), 0, bold, cell)
        if bold:
            words = len(text.split())
            suffix = o["button_suffix"]
            if suffix and text.endswith(suffix):
                return Node(uid, "button", text[: -len(suffix)].strip(), 0, True, cell)
            if words <= o["button_max_words"]:
                return Node(uid, "button", text, 0, True, cell)
            return Node(uid, "label", text, 0, True, cell)
        return Node(uid, "paragraph", text, 0, False, cell)

    def walk_tbl(tbl, out):
        for tr in tbl.findall(qn("w:tr")):
            for tc in tr.findall(qn("w:tc")):
                counters["cell"] += 1
                walk(tc, counters["cell"], out)

    def walk(parent, cell, out):
        for ch in parent:
            if ch.tag == qn("w:p"):
                n = make_node(ch, cell)
                if n:
                    out.append(n)
            elif ch.tag == qn("w:tbl"):
                walk_tbl(ch, out)
            elif ch.tag == qn("w:sdt"):
                c = ch.find(qn("w:sdtContent"))
                if c is not None:
                    walk(c, cell, out)

    idx = 0
    for ch in document.element.body:
        if ch.tag == qn("w:tbl"):
            nodes: List[Node] = []
            walk_tbl(ch, nodes)
            doc.sections.append(Section(idx, f"table {idx}", nodes))
            idx += 1
        elif ch.tag == qn("w:p"):
            counters["cell"] += 1
            nodes = []
            n = make_node(ch, counters["cell"])
            if not n:
                continue
            if n.kind == "heading" and n.level == 1 and not doc.title:
                doc.title = n.text
                continue
            doc.sections.append(Section(idx, f"paragraph {idx}", [n]))
            idx += 1
    for s in document.sections:
        for part in (s.header, s.footer):
            for p in part.paragraphs:
                if p.text.strip():
                    doc.notes.append(p.text.strip())
    for sec in doc.sections:
        derive(sec)
    return doc


def derive(sec: Section):
    """Turn a flat, ordered node stream into title / intro / cards / buttons."""
    nodes = sec.nodes
    title = next((n for n in nodes if n.kind == "heading"), None) or next(
        (n for n in nodes if n.kind == "label"), None
    )
    sec.title = title
    sec.placeholders = [n for n in nodes if n.kind == "placeholder"]
    sec.items = [n for n in nodes if n.kind == "listitem"]
    in_card = set()
    cards: List[Card] = []
    if title is not None and title.kind == "heading":
        start = nodes.index(title) + 1
        i = start
        while i < len(nodes):
            n = nodes[i]
            if n.kind == "heading" and n.level > title.level:
                c = Card(n)
                if i > 0 and nodes[i - 1].kind == "placeholder" and nodes[i - 1].cell == n.cell:
                    c.pre_placeholders.append(nodes[i - 1])
                    in_card.add(nodes[i - 1].uid)
                j = i + 1
                while j < len(nodes) and nodes[j].cell == n.cell and not (
                    nodes[j].kind == "heading" and nodes[j].level > title.level
                ):
                    m = nodes[j]
                    {"button": c.buttons, "placeholder": c.placeholders}.get(m.kind, c.body).append(m)
                    in_card.add(m.uid)
                    j += 1
                in_card.add(n.uid)
                cards.append(c)
                i = j
            else:
                i += 1
    sec.cards = cards
    first_card_pos = min((nodes.index(c.heading) for c in cards), default=len(nodes))
    tpos = nodes.index(title) if title else -1
    for pos, n in enumerate(nodes):
        if n is title or n.uid in in_card or n.kind != "paragraph":
            continue
        (sec.intro if pos < first_card_pos and pos > tpos else sec.outro).append(n)
    sec.buttons = [n for n in nodes if n.kind == "button" and n.uid not in in_card]


# ----------------------------------------------------------------------------
# template model (Elementor)
# ----------------------------------------------------------------------------
TEXT_FIELDS = {
    "heading": ["title"],
    "text-editor": ["editor"],
    "button": ["text"],
    "call-to-action": ["title", "description", "button"],
    "image-box": ["title_text", "description_text"],
    "shortcode": ["shortcode"],
    "nested-accordion": ["items"],
}
DEFAULT_AS = {
    ("heading", "title"): "heading",
    ("text-editor", "editor"): "paragraphs",
}


class Template:
    def __init__(self, data):
        self.data = data
        self.idx = {}  # id -> dict(el, path, parent)
        for i, el in enumerate(data.get("content", [])):
            self._index(el, f"content[{i}]", None)

    def _index(self, el, path, parent):
        if el["id"] in self.idx:
            raise ValueError(f"duplicate element id in template: {el['id']}")
        self.idx[el["id"]] = {"el": el, "path": path, "parent": parent}
        for i, c in enumerate(el.get("elements", [])):
            self._index(c, f"{path}.elements[{i}]", el["id"])

    def walk(self):
        for k, v in list(self.idx.items()):
            yield k, v

    def new_id(self):
        while True:
            i = secrets.token_hex(4)
            if i not in self.idx:
                return i


# ----------------------------------------------------------------------------
# mapping engine
# ----------------------------------------------------------------------------
class Engine:
    def __init__(self, doc: SourceDoc, tpl_data: dict, cfg: dict):
        self.doc = doc
        self.orig = copy.deepcopy(tpl_data)
        self.tpl = Template(copy.deepcopy(tpl_data))
        self.orig_tpl = Template(copy.deepcopy(tpl_data))
        self.cfg = cfg
        self.log: List[dict] = []  # mapping records
        self.warnings: List[str] = []
        self.errors: List[str] = []
        self.clones: List[dict] = []
        self.written = set()  # (target id, field)
        self.vars = {}

    # ---- vars / expressions ------------------------------------------------
    def compute_vars(self):
        for name, spec in (self.cfg.get("vars") or {}).items():
            src = self.doc.title if spec.get("from") == "doc.title" else ""
            m = re.search(spec["regex"], src)
            if not m:
                self.warnings.append(
                    f"VAR '{name}': regex {spec['regex']!r} did not match document title {src!r}; dependent fields left unchanged"
                )
                continue
            self.vars[name] = m.group(spec.get("group", 1)).strip()
            self.warnings.append(
                f"DERIVED: var '{name}' = {self.vars[name]!r} derived from document title {src!r} (not an explicit content element) - please confirm"
            )

    def resolve(self, expr, sec: Optional[Section], item=None):
        m = re.match(r"^([A-Za-z_.]+?)(?:\[(\d+)\])?$", expr)
        if not m:
            raise ValueError(f"bad value expression {expr!r}")
        root, i = m.group(1), m.group(2)
        if root.startswith("var."):
            return self.vars.get(root[4:])
        if root == "doc.title":
            return self.doc.title
        if root.startswith("item."):
            part = root[5:]
            if isinstance(item, Card):
                v = {"heading": item.heading, "body": item.body,
                     "button": item.buttons[0] if item.buttons else None}.get(part)
            else:
                v = item if part == "text" else None
            return v[int(i)] if i is not None and isinstance(v, list) and int(i) < len(v) else (None if i is not None else v)
        if sec is None:
            raise ValueError(f"expression {expr!r} needs a section")
        v = {"title": sec.title, "paras": sec.intro, "outro": sec.outro, "buttons": sec.buttons,
             "cards": sec.cards, "items": sec.items}.get(root)
        if i is not None:
            return v[int(i)] if isinstance(v, list) and int(i) < len(v) else None
        return v

    # ---- serialisation -----------------------------------------------------
    @staticmethod
    def nodes_of(v):
        if v is None or isinstance(v, str):
            return []
        if isinstance(v, Node):
            return [v]
        if isinstance(v, Card):
            return [v.heading] + v.body + v.buttons
        return [n for x in v for n in Engine.nodes_of(x)]

    @staticmethod
    def text_of(v):
        if isinstance(v, str):
            return v
        return " ".join(n.text.replace("\n", " ") for n in Engine.nodes_of(v))

    def serialise(self, v, how, old, bind, sec):
        if how == "heading":
            new = esc(self.text_of(v))
            if norm(new) == norm(old):  # identical wording -> keep template markup (e.g. <span>)
                return old
            span = bind.get("span_last_words")
            if span:
                w = new.split()
                if len(w) > span:
                    new = " ".join(w[:-span]) + " <span>" + " ".join(w[-span:]) + "</span>"
            return new
        if how == "paragraphs":
            nodes = self.nodes_of(v)
            return "".join(f"<p>{esc(n.text).replace(chr(10), '<br>')}</p>" for n in nodes)
        if how == "compose":
            out = []
            for part in bind["compose"]:
                if "paragraphs" in part:
                    out.append(self.serialise(self.resolve(part["paragraphs"], sec), "paragraphs", old, bind, sec))
                elif "list" in part:
                    lis = []
                    for c in self.resolve(part["list"], sec) or []:
                        lis.append("<li>" + part["item"].format(
                            heading=esc(c.heading.text), body=esc(" ".join(b.text for b in c.body))) + "</li>")
                    out.append("<ul>" + "".join(lis) + "</ul>")
                if "after" in part:
                    out.append(self.serialise(self.resolve(part["after"], sec), "paragraphs", old, bind, sec))
            return "".join(out)
        return self.text_of(v)  # "text": raw

    # ---- writing -----------------------------------------------------------
    def write(self, rule, sec, target_id, fld, bind, item=None, value=None):
        ent = self.tpl.idx.get(target_id)
        if not ent:
            self.errors.append(f"[{rule}] target id {target_id} not found in template")
            return
        el = ent["el"]
        st = el["settings"]
        if fld not in st:
            self.errors.append(f"[{rule}] field '{fld}' does not exist on {target_id} ({el.get('widgetType')}); not adding new keys")
            return
        if (target_id, fld) in self.written:
            self.errors.append(f"[{rule}] DUPLICATE write to {target_id}.{fld}")
            return
        how = bind.get("as") or ("compose" if "compose" in bind else DEFAULT_AS.get((el.get("widgetType"), fld), "text"))
        if how == "compose":
            v = None
        else:
            v = value if value is not None else self.resolve(bind["value"], sec, item)
            if v is None or (isinstance(v, (list, str)) and not v):
                self.warnings.append(f"[{rule}] no source value for {target_id}.{fld} (expr {bind.get('value')!r}); template value kept")
                return
        old = st[fld]
        new = self.serialise(v, how, old, bind, sec)
        if type(new) is not type(old):
            self.errors.append(f"[{rule}] type mismatch on {target_id}.{fld}")
            return
        if how == "heading" and "<span" in str(old) and new != old and "<span" not in new:
            self.warnings.append(
                f"[{rule}] heading {target_id}: template had an accent <span> ({excerpt(old, 60)!r}); source gives plain text so the accent was dropped "
                f"(set 'span_last_words' on the bind to re-apply)")
        nodes = self.nodes_of(v) if how != "compose" else (
            [n for p in bind["compose"] for k in ("paragraphs", "list", "after") if k in p
             for n in self.nodes_of(self.resolve(p[k], sec))])
        for n in nodes:
            n.used.append((rule, target_id, fld))
        st[fld] = new
        self.written.add((target_id, fld))
        self.log.append({
            "rule": rule, "section": sec.index if sec else None,
            "source": " | ".join(n.short() for n in nodes[:3]) + (" …" if len(nodes) > 3 else "") if nodes else f"(derived) {v}",
            "source_kind": ",".join(sorted({n.kind for n in nodes})) or "derived",
            "target": target_id, "path": ent["path"], "widget": el.get("widgetType"), "field": fld,
            "old": excerpt(old, 50), "new": excerpt(new, 70),
            "status": "unchanged (identical wording)" if new == old else "mapped",
            "_nodes": [n.uid for n in nodes], "_new": new,
        })

    def clone(self, rule, el, parent_el):
        c = copy.deepcopy(el)

        def renew(e):
            old = e["id"]
            e["id"] = self.tpl.new_id()
            self.tpl.idx[e["id"]] = {"el": e, "path": "(clone)", "parent": parent_el["id"] if e is c else None}
            self.clones.append({"rule": rule, "from": old, "id": e["id"]})
            for x in e.get("elements", []):
                renew(x)

        renew(c)
        parent_el["elements"].append(c)
        return c

    # ---- bind types -------------------------------------------------------
    def do_repeat(self, rule, sec, b):
        rp = b["repeat"]
        items = self.resolve(rp["over"], sec)
        ent = self.tpl.idx.get(rp["container"])
        if not ent:
            self.errors.append(f"[{rule}] repeat container {rp['container']} not found")
            return
        cont = ent["el"]
        slots = cont["elements"]
        items = items or []
        pol = {**self.cfg.get("defaults", {}), **rp}
        if len(items) > len(slots):
            if pol.get("on_overflow", "clone") == "clone":
                src_i = pol.get("clone_from", len(slots) - 1)
                base = slots[src_i]
                n_new = len(items) - len(slots)
                for _ in range(n_new):
                    c = self.clone(rule, base, cont)
                self.warnings.append(
                    f"[{rule}] source has {len(items)} items but template container {rp['container']} has {len(slots) - n_new}; "
                    f"cloned slot {base['id']} x{n_new} (new keys follow the cloned widget; review link/style)")
            else:
                self.warnings.append(f"[{rule}] {len(items) - len(slots)} item(s) have no template slot and were NOT placed")
        elif len(items) < len(slots):
            self.warnings.append(
                f"[{rule}] template container {rp['container']} has {len(slots) - len(items)} more slot(s) than source items; "
                f"surplus slots keep template text (see retained list)")
        for slot, it in zip(cont["elements"], items):
            for sb in b["binds"]:
                self.write(rule, sec, slot["id"], sb["field"], sb, item=it)

    def do_accordion(self, rule, sec, b):
        a = b["accordion"]
        ent = self.tpl.idx.get(a["target"])
        if not ent:
            self.errors.append(f"[{rule}] accordion {a['target']} not found")
            return
        w = ent["el"]
        items = self.resolve(a["over"], sec) or []
        lst, kids = w["settings"]["items"], w["elements"]
        if len(lst) != len(kids):
            self.errors.append(f"[{rule}] accordion items ({len(lst)}) / child containers ({len(kids)}) out of sync in template")
            return
        n_tpl = len(lst)
        while len(lst) < len(items):
            base_i = len(lst) - 1
            new_item = copy.deepcopy(lst[base_i])
            new_item["_id"] = secrets.token_hex(4)[:7]
            lst.append(new_item)
            self.clone(rule, kids[base_i], w)
        if len(items) > n_tpl:
            self.warnings.append(f"[{rule}] accordion had {n_tpl} items, source has {len(items)}; cloned last item/container x{len(items) - n_tpl}")
        elif len(items) < n_tpl:
            self.warnings.append(f"[{rule}] accordion has {n_tpl - len(items)} more item(s) than source; surplus keep template text")
        for i, c in enumerate(items):
            q = self.resolve(a["question"], sec, c)
            lst[i]["item_title"] = self.text_of(q)
            for n in self.nodes_of(q):
                n.used.append((rule, a["target"], f"items[{i}].item_title"))
            ed = None
            stack = list(w["elements"][i].get("elements", []))
            while stack:
                e = stack.pop(0)
                if e.get("widgetType") == "text-editor":
                    ed = e
                    break
                stack = e.get("elements", []) + stack
            if ed is None:
                self.errors.append(f"[{rule}] no text-editor in accordion child {i}")
                continue
            ans = self.resolve(a["answer"], sec, c)
            ed["settings"]["editor"] = self.serialise(ans, "paragraphs", "", {}, sec)
            for n in self.nodes_of(ans):
                n.used.append((rule, ed["id"], "editor"))
            self.written.add((ed["id"], "editor"))
            self.log.append({
                "rule": rule, "section": sec.index, "source": f"Q: {excerpt(q.text, 45)} | A: {excerpt(self.text_of(ans), 45)}",
                "source_kind": "heading,paragraph", "target": a["target"], "path": ent["path"],
                "widget": "nested-accordion", "field": f"items[{i}] + child editor {ed['id']}",
                "old": "", "new": excerpt(lst[i]["item_title"], 60), "status": "mapped",
                "_nodes": [q.uid] + [n.uid for n in self.nodes_of(ans)],
                "_new": lst[i]["item_title"] + " " + ed["settings"]["editor"],
            })
        self.written.add((a["target"], "items"))

    # ---- sections -----------------------------------------------------------
    def find_section(self, spec):
        if "index" in spec:
            for s in self.doc.sections:
                if s.index == spec["index"]:
                    return s
            return None
        rx = re.compile(spec["title_regex"], re.I)
        return next((s for s in self.doc.sections if s.title and rx.search(s.title.text)), None)

    @staticmethod
    def check_expect(val, exp):
        m = re.match(r"^(>=|<=|==|>|<)?\s*(\d+)$", str(exp))
        op, n = (m.group(1) or "=="), int(m.group(2))
        return {">=": val >= n, "<=": val <= n, "==": val == n, ">": val > n, "<": val < n}[op]

    def run(self):
        self.compute_vars()
        for ig in self.cfg.get("ignore", []):
            for s in self.doc.sections:
                if s.index == ig["section"]:
                    s.ignored = ig["reason"]
        for rule in self.cfg["rules"]:
            name = rule["name"]
            sec = None
            if "section" in rule:
                sec = self.find_section(rule["section"])
                if sec is None:
                    self.errors.append(f"[{name}] source section {rule['section']} not found - rule skipped")
                    continue
                sec.rule = (sec.rule + "," + name) if sec.rule else name
                for k, exp in (rule["section"].get("expect") or {}).items():
                    cur = {"cards": len(sec.cards), "paras": len(sec.intro), "buttons": len(sec.buttons),
                           "items": len(sec.items)}[k]
                    if not self.check_expect(cur, exp):
                        self.warnings.append(f"[{name}] STRUCTURE DRIFT: section {sec.index} has {k}={cur}, config expected {exp}")
            for b in rule["binds"]:
                if "repeat" in b:
                    self.do_repeat(name, sec, b)
                elif "accordion" in b:
                    self.do_accordion(name, sec, b)
                else:
                    self.write(name, sec, b["target"], b["field"], b)
        pg = self.cfg.get("page", {})
        if "title" in pg:
            try:
                new = pg["title"].format(**{f"var_{k}": v for k, v in self.vars.items()}) if "var_" in pg["title"] else pg["title"]
                new = re.sub(r"\{var\.(\w+)\}", lambda m: self.vars[m.group(1)], pg["title"])
                self.tpl.data["title"] = new
                self.log.append({"rule": "page", "section": None, "source": "(derived) document title", "source_kind": "derived",
                                 "target": "(page)", "path": "title", "widget": "page", "field": "title",
                                 "old": self.orig["title"], "new": new, "status": "mapped", "_nodes": [], "_new": new})
            except KeyError as e:
                self.warnings.append(f"[page] var {e} unavailable; page title unchanged")


# ----------------------------------------------------------------------------
# validation
# ----------------------------------------------------------------------------
def validate(eng: Engine, out: dict):
    errs, warns = list(eng.errors), []
    checks = []

    def chk(name, ok, detail=""):
        checks.append({"check": name, "ok": ok, "detail": detail})
        if not ok:
            errs.append(f"VALIDATION: {name} - {detail}")

    # 1 syntax
    try:
        json.loads(json.dumps(out, ensure_ascii=False))
        chk("JSON serialises and re-parses", True)
    except Exception as e:  # pragma: no cover
        chk("JSON serialises and re-parses", False, str(e))

    # 2 structure preserved
    def smap(data):
        m = {}

        def rec(e, parent):
            m[e["id"]] = (e.get("elType"), e.get("widgetType"), {k: type(v).__name__ for k, v in e["settings"].items()},
                          parent, [c["id"] for c in e.get("elements", [])])
            for c in e.get("elements", []):
                rec(c, e["id"])
        for e in data["content"]:
            rec(e, None)
        return m

    a, b = smap(eng.orig), smap(out)
    clone_ids = {c["id"] for c in eng.clones}
    problems = []
    if set(out) != set(eng.orig):
        problems.append("top-level keys differ")
    for k in ("version", "type"):
        if out.get(k) != eng.orig.get(k):
            problems.append(f"top-level '{k}' changed")
    if out.get("page_settings") != eng.orig.get("page_settings"):
        problems.append("page_settings changed")
    for i, v in a.items():
        if i not in b:
            problems.append(f"element {i} missing")
            continue
        w = b[i]
        if v[0] != w[0] or v[1] != w[1]:
            problems.append(f"{i}: elType/widgetType changed")
        if v[2] != w[2]:
            problems.append(f"{i}: settings keys/types changed")
        if v[3] != w[3]:
            problems.append(f"{i}: parent changed")
        if [x for x in w[4] if x not in clone_ids] != v[4]:
            problems.append(f"{i}: original children reordered/removed")
    extra = set(b) - set(a)
    if extra - clone_ids:
        problems.append(f"unexpected new ids {sorted(extra - clone_ids)}")
    chk("Template structure, key names and data types preserved", not problems, "; ".join(problems[:5]))
    chk("New elements exist only as logged clones", not (extra - clone_ids), f"{len(clone_ids)} clone(s) logged")

    # 3 ids unique
    ids = []

    def collect(e):
        ids.append(e["id"])
        for c in e.get("elements", []):
            collect(c)
    for e in out["content"]:
        collect(e)
    chk("Element ids unique", len(ids) == len(set(ids)), f"{len(ids)} elements")

    # 4 accordion consistency
    acc = []

    def walk(e):
        if e.get("widgetType") == "nested-accordion":
            acc.append(len(e["settings"]["items"]) == len(e["elements"]))
        for c in e.get("elements", []):
            walk(c)
    for e in out["content"]:
        walk(e)
    chk("Accordion items match child containers", all(acc))

    # 5 coverage
    unmapped_sections, unmapped_nodes, dup = [], [], []
    for s in eng.doc.sections:
        if s.ignored:
            continue
        if not s.rule:
            unmapped_sections.append(s)
            continue
        for n in s.nodes:
            if n.kind in CONTENT_KINDS and not n.used:
                unmapped_nodes.append((s, n))
            if len({(u[1], u[2]) for u in n.used}) > 1:
                dup.append((s, n))
    chk("No relevant source node silently skipped", True,
        f"{len(unmapped_sections)} unmapped section(s), {len(unmapped_nodes)} unmapped node(s) in mapped sections - all reported")
    chk("No source node mapped to two different targets", not dup,
        "; ".join(f"{n.short()}" for _, n in dup[:3]))

    # 6 text fidelity
    flat = {}

    def fl(e):
        for k, v in e["settings"].items():
            flat[(e["id"], k)] = v if isinstance(v, str) else json.dumps(v, ensure_ascii=False)
        for c in e.get("elements", []):
            fl(c)
    for e in out["content"]:
        fl(e)
    all_text = norm(" ".join(strip_tags(v) for v in flat.values()))
    by_uid = {n.uid: n for s in eng.doc.sections for n in s.nodes}
    bad = []
    for r in eng.log:
        for uid in r["_nodes"]:
            n = by_uid[uid]
            if norm(esc(n.text)) not in norm(r["_new"]) and norm(n.text) not in norm(r["_new"]):
                bad.append(n.short())
    chk("Every mapped text appears verbatim in its target field", not bad, "; ".join(bad[:3]))
    chk("Source wording unmodified (spot check on whole output)", all(
        norm(n.text) in all_text for s in eng.doc.sections for n in s.nodes if n.used), "")

    # 7 retained-from-template
    stale = set()
    old_t = eng.orig.get("title", "")
    new_t = out.get("title", "")
    for w in re.findall(r"[a-z]{4,}", old_t.lower()):
        if w not in new_t.lower():
            stale.add(w)
    stale |= {x.lower() for x in eng.cfg.get("stale_terms", [])}
    retained, media = [], []
    reasons = {r["target"]: r["reason"] for r in eng.cfg.get("retain", [])}
    for i, ent in eng.tpl.idx.items():
        el = ent["el"]
        wt = el.get("widgetType")
        for f in TEXT_FIELDS.get(wt, []):
            if (i, f) in eng.written:
                continue
            val = el["settings"].get(f)
            if f == "items":
                continue
            if isinstance(val, str) and norm(val):
                hit = sorted(t for t in stale if t in norm(val).lower())
                retained.append({"id": i, "path": ent["path"], "widget": wt, "field": f, "value": excerpt(val, 90),
                                 "reason": reasons.get(i, "no mapping rule"), "stale": hit})
        if wt == "nested-accordion" and (i, "items") not in eng.written:
            retained.append({"id": i, "path": ent["path"], "widget": wt, "field": "items", "value": "(accordion untouched)",
                             "reason": reasons.get(i, "no mapping rule"), "stale": []})
        if wt == "template":
            retained.append({"id": i, "path": ent["path"], "widget": wt, "field": "template_id",
                             "value": f"global template #{el['settings'].get('template_id')}",
                             "reason": reasons.get(i, "global Elementor template - content lives outside this JSON"), "stale": []})
        if wt in ("image", "call-to-action", "image-box"):
            for k in ("image", "graphic_image"):
                v = el["settings"].get(k)
                if isinstance(v, dict):
                    u = v.get("url", "") or ""
                    media.append({"id": i, "path": ent["path"], "widget": wt, "url": u or "(empty)", "alt": v.get("alt", ""),
                                  "stale": sorted(t for t in stale if t in u.lower() or t in (v.get("alt", "") or "").lower())})
    for m in media:
        if m["stale"]:
            warns.append(f"STALE TEMPLATE MEDIA retained in {m['id']} ({m['widget']}): file name contains {m['stale']} - replace image")
    for r in retained:
        if r["stale"]:
            warns.append(f"STALE TEMPLATE TEXT retained in {r['id']} ({r['widget']}.{r['field']}): {r['value']!r} contains {r['stale']}")
    chk("No stale template terms left in text fields", not any(r["stale"] for r in retained),
        f"terms checked: {sorted(stale)}")

    return {
        "errors": errs, "warnings": eng.warnings + warns, "checks": checks,
        "unmapped_sections": unmapped_sections, "unmapped_nodes": unmapped_nodes,
        "retained": retained, "media": media,
    }


# ----------------------------------------------------------------------------
# reporting
# ----------------------------------------------------------------------------
def sec_text(s: Section):
    return [f"{n.kind}{n.level or ''}: {n.text}" for n in s.nodes]


def build_report(eng: Engine, res: dict, args) -> str:
    L = []
    mapped = [r for r in eng.log if r["status"] == "mapped"]
    L.append("# Content → JSON mapping report\n")
    L.append(f"- Content file: `{args.content}`  \n- Template: `{args.template}`  \n- Mapping config: `{args.mapping}`  \n- Output: `{args.out}`\n")
    L.append(f"**Result:** {'FAILED' if res['errors'] else 'PASSED'} · {len(mapped)} fields written · "
             f"{sum(1 for r in eng.log if r['status'].startswith('unchanged'))} unchanged (identical wording) · "
             f"{len(eng.clones)} cloned element(s) · {len(res['warnings'])} warning(s) · {len(res['errors'])} error(s)\n")
    L.append("## 1. Mapping summary\n")
    L.append("| Content section | Rule | Source content | → Template element | Field | Status |")
    L.append("|---|---|---|---|---|---|")
    for r in eng.log:
        sec = "-" if r["section"] is None else r["section"]
        L.append(f"| {sec} | {r['rule']} | {r['source'].replace('|', '/')} | `{r['target']}` ({r['widget']}) `{r['path']}` | {r['field']} | {r['status']} |")
    L.append("\n## 2. Unmatched source content (no template element)\n")
    if not res["unmapped_sections"] and not res["unmapped_nodes"]:
        L.append("None.\n")
    for s in res["unmapped_sections"]:
        L.append(f"**Source section {s.index} ({s.label})** - no rule maps it; template has no clear counterpart:")
        for t in sec_text(s):
            L.append(f"- {t}")
        L.append("")
    for s, n in res["unmapped_nodes"]:
        L.append(f"- Section {s.index}: `{n.kind}` not placed anywhere: {n.text!r}")
    ig = [s for s in eng.doc.sections if s.ignored]
    L.append("\n### Intentionally ignored (by config)\n")
    for s in ig:
        L.append(f"- Section {s.index} ({s.label}): {s.ignored}  — text: " + "; ".join(excerpt(n.text, 30) for n in s.nodes[:12]))
    L.append("\n### Source placeholders (instructions, no content to map)\n")
    for s in eng.doc.sections:
        in_cards = {x.uid for c in s.cards for x in c.pre_placeholders + c.placeholders}
        for n in s.placeholders:
            if n.uid not in in_cards:
                L.append(f"- Section {s.index}: {n.text!r}")
        for c in s.cards:
            for n in c.pre_placeholders + c.placeholders:
                L.append(f"- Section {s.index} (card '{c.heading.short()}'): {n.text!r}")
    if eng.doc.notes:
        L.append("\n### Document notes (header/footer text, not page content)\n")
        for n in eng.doc.notes:
            L.append(f"- {n}")
    L.append("\n## 3. Template fields retained (no source value)\n")
    L.append("| Element | Widget.field | Current value | Reason |\n|---|---|---|---|")
    for r in res["retained"]:
        flag = " ⚠ stale: " + ",".join(r["stale"]) if r["stale"] else ""
        L.append(f"| `{r['id']}` | {r['widget']}.{r['field']} | {r['value'].replace('|', '/')}{flag} | {r['reason']} |")
    L.append("\n### Media retained from template (source only has placeholders)\n")
    for m in res["media"]:
        L.append(f"- `{m['id']}` {m['widget']} `{m['path']}` url={m['url']} alt={m['alt']!r}" + (f" ⚠ stale: {','.join(m['stale'])}" if m["stale"] else ""))
    L.append("\n## 4. Warnings / ambiguities\n")
    for w in res["warnings"]:
        L.append(f"- {w}")
    if not res["warnings"]:
        L.append("None.")
    L.append("\n## 5. Validation\n")
    for c in res["checks"]:
        L.append(f"- {'✅' if c['ok'] else '❌'} {c['check']}" + (f" - {c['detail']}" if c["detail"] else ""))
    for e in res["errors"]:
        L.append(f"- ❌ {e}")
    if eng.clones:
        L.append("\n### Cloned elements\n")
        for c in eng.clones:
            L.append(f"- `{c['id']}` cloned from `{c['from']}` (rule {c['rule']})")
    return "\n".join(L) + "\n"


# ----------------------------------------------------------------------------
# inspect
# ----------------------------------------------------------------------------
def cmd_inspect(args):
    cfg = json.load(open(args.mapping)) if args.mapping else {}
    doc = parse_docx(args.content, (cfg.get("source") or {}).get("options"))
    print(f"DOCUMENT TITLE: {doc.title!r}\n")
    for s in doc.sections:
        print(f"[{s.index}] {s.label}: title={s.title.short() if s.title else None!r} "
              f"intro={len(s.intro)} cards={len(s.cards)} buttons={[b.text for b in s.buttons]} "
              f"items={len(s.items)} placeholders={len(s.placeholders)}")
        for c in s.cards[:3]:
            print(f"      card: {c.heading.short()!r} body={len(c.body)} btn={[b.text for b in c.buttons]}")
    print("\nTEMPLATE WIDGETS")
    t = Template(json.load(open(args.template)))
    for i, ent in t.walk():
        el = ent["el"]
        if el.get("elType") == "widget":
            f = TEXT_FIELDS.get(el["widgetType"], [])
            vals = {k: excerpt(el["settings"][k], 45) for k in f if k in el["settings"] and k != "items"}
            print(f"  {i:>9} {el['widgetType']:<16} {ent['path']:<45} {vals}")
        else:
            print(f"  {i:>9} container        {ent['path']}")


def run_mapping(docx_src, tpl_data, cfg):
    """docx_src: path or file-like. Returns (engine, validation_result)."""
    doc = parse_docx(docx_src, (cfg.get("source") or {}).get("options"))
    eng = Engine(doc, tpl_data, cfg)
    eng.run()
    return eng, validate(eng, eng.tpl.data)


def cmd_map(args):
    cfg = json.load(open(args.mapping, encoding="utf8"))
    tpl = json.load(open(args.template, encoding="utf8"))
    eng, res = run_mapping(args.content, tpl, cfg)
    out = eng.tpl.data
    json.dump(out, open(args.out, "w", encoding="utf8"), ensure_ascii=False, indent=2)
    rep = build_report(eng, res, args)
    if args.report:
        open(args.report, "w", encoding="utf8").write(rep)
        slim = {k: v for k, v in res.items() if k not in ("unmapped_sections", "unmapped_nodes")}
        slim["unmapped_sections"] = [{"index": s.index, "text": sec_text(s)} for s in res["unmapped_sections"]]
        slim["unmapped_nodes"] = [{"section": s.index, "kind": n.kind, "text": n.text} for s, n in res["unmapped_nodes"]]
        slim["mapping"] = [{k: v for k, v in r.items() if not k.startswith("_")} for r in eng.log]
        slim["clones"] = eng.clones
        json.dump(slim, open(args.report.rsplit(".", 1)[0] + ".json", "w", encoding="utf8"), ensure_ascii=False, indent=2)
    print(rep)
    bad = bool(res["errors"]) or (args.strict and (res["warnings"] or res["unmapped_sections"] or res["unmapped_nodes"]))
    sys.exit(1 if bad else 0)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest="cmd", required=True)
    for name in ("inspect", "map"):
        p = sp.add_parser(name)
        p.add_argument("--content", required=True)
        p.add_argument("--template", required=True)
        p.add_argument("--mapping", required=(name == "map"))
    sp.choices["map"].add_argument("--out", required=True)
    sp.choices["map"].add_argument("--report")
    sp.choices["map"].add_argument("--strict", action="store_true", help="non-zero exit on any warning/unmapped content")
    a = ap.parse_args()
    cmd_inspect(a) if a.cmd == "inspect" else cmd_map(a)


if __name__ == "__main__":
    main()
